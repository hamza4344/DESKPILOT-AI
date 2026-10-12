"""Modified for DeskPilot AI: PDF/text readers and DOCX table extraction."""

from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return [
            Document(
                page_content=page.extract_text() or "",
                metadata={"source": str(path), "page": number},
            )
            for number, page in enumerate(PdfReader(path).pages)
        ]
    if suffix in {".txt", ".md"}:
        return [
            Document(page_content=path.read_text(encoding="utf-8"), metadata={"source": str(path)})
        ]
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        rows = [
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        ]
        return [
            Document(page_content="\n".join([*paragraphs, *rows]), metadata={"source": str(path)})
        ]
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900, chunk_overlap=120, add_start_index=True
    )
    return splitter.split_documents(list(docs))
