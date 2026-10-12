from docx import Document

from app.services.ingestion import chunk_documents, load_file


def test_docx_tables_are_searchable(tmp_path):
    path = tmp_path / "guide.docx"
    doc = Document()
    doc.add_paragraph("Support directory")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "VPN"
    row.cells[1].text = "Service desk"
    doc.save(path)
    chunks = chunk_documents(load_file(path))
    assert "VPN | Service desk" in chunks[0].page_content
